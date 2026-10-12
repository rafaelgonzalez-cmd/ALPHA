import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── Denomination mapping ──────────────────────────────────────────────────────
DEN_MAP = {
    "d_1000": 1000,
    "d_500": 500,
    "d_200": 200,
    "d_100": 100,
    "d_50": 50,
    "d_20": 20,
    "d_10": 10,
    "d_5": 5,
    "d_2": 2,
    "d_1": 1,
    "d_050": 0.50,
}


def format_currency(value: float) -> str:
    return f"${value:,.2f}"


def calc_total_from_denominaciones(denominaciones: dict) -> float:
    total = 0.0
    for col, unit_value in DEN_MAP.items():
        total += denominaciones.get(col, 0) * unit_value
    return round(total, 2)


# ── Excel export ──────────────────────────────────────────────────────────────

def _thin_border() -> Border:
    s = Side(style="thin")
    return Border(left=s, right=s, top=s, bottom=s)


def _apply_header(cell, text: str):
    cell.value = text
    cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    cell.font = Font(color="FFFFFF", bold=True, size=11)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = _thin_border()


def export_estado_cuenta_excel(estado_cuenta: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Estado de Cuenta"

    headers = [
        "FECHA",
        "SALDO INICIAL",
        "INGRESOS EN EFECTIVO",
        "EGRESOS EN EFECTIVO",
        "SALDO FINAL",
    ]

    # Header row
    ws.row_dimensions[1].height = 28
    for col_idx, h in enumerate(headers, 1):
        _apply_header(ws.cell(row=1, column=col_idx), h)

    # Data rows
    even_fill = PatternFill(start_color="DDEEFF", end_color="DDEEFF", fill_type="solid")
    odd_fill  = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    currency_fmt = '"$"#,##0.00'

    for row_idx, row in enumerate(estado_cuenta, 2):
        fill = even_fill if row_idx % 2 == 0 else odd_fill
        values = [
            row["FECHA"],
            row["SALDO INICIAL"],
            row["INGRESOS EN EFECTIVO"],
            row["EGRESOS EN EFECTIVO"],
            row["SALDO FINAL"],
        ]
        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = _thin_border()
            cell.fill = fill
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.number_format = currency_fmt
                cell.alignment = Alignment(horizontal="right")

    # Totals row
    if estado_cuenta:
        total_row = len(estado_cuenta) + 2
        total_fill = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
        bold = Font(bold=True, size=11)

        for col_idx in range(1, 6):
            cell = ws.cell(row=total_row, column=col_idx)
            cell.border = _thin_border()
            cell.fill = total_fill
            cell.font = bold

        ws.cell(row=total_row, column=1).value = "TOTALES"
        ws.cell(row=total_row, column=1).alignment = Alignment(horizontal="center")

        ws.cell(row=total_row, column=3).value = sum(r["INGRESOS EN EFECTIVO"] for r in estado_cuenta)
        ws.cell(row=total_row, column=3).number_format = currency_fmt
        ws.cell(row=total_row, column=3).alignment = Alignment(horizontal="right")

        ws.cell(row=total_row, column=4).value = sum(r["EGRESOS EN EFECTIVO"] for r in estado_cuenta)
        ws.cell(row=total_row, column=4).number_format = currency_fmt
        ws.cell(row=total_row, column=4).alignment = Alignment(horizontal="right")

        # Saldo final del último día
        ws.cell(row=total_row, column=5).value = estado_cuenta[-1]["SALDO FINAL"]
        ws.cell(row=total_row, column=5).number_format = currency_fmt
        ws.cell(row=total_row, column=5).alignment = Alignment(horizontal="right")

    # Column widths
    for i, w in enumerate([14, 18, 22, 22, 18], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Freeze header
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
